from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\nnjj1\UMD-work\dairy_protein_USDA")
OUT = ROOT / "report_summary" / "ACS2026" / "ACS2026_speaker_script_YW_0823.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
RED = "E21836"
GOLD = "FFD200"
GRAY = "5A5A5A"
LIGHT = "F4F6F9"
PALE_GOLD = "FFF8E1"
WHITE = "FFFFFF"


SLIDES = [
    {
        "n": 1,
        "title": "Title",
        "seconds": 35,
        "screen": "Talk title, authors, University of Maryland affiliation, and ACS meeting line.",
        "script": [
            "Good afternoon. My name is Yi Wang, from the Department of Nutrition and Food Science at the University of Maryland. It is a pleasure to be here in Chicago.",
            "Today I will talk about what happens when a zein nanoparticle meets dairy proteins. The key idea is that the particle does not keep the clean surface we designed in the lab. In a real food, proteins quickly coat it and give it a new identity - what we call a protein corona. We are using computational modeling to understand and eventually predict that change."
        ],
        "cue": "Open conversationally. Do not read the full title. Emphasize 'new identity' and then name the protein corona.",
    },
    {
        "n": 2,
        "title": "Outline",
        "seconds": 30,
        "screen": "Five-part outline: context, DLVO, docking, coarse-grained MD, and significance.",
        "script": [
            "The talk follows three questions, moving from a whole-particle view to a molecular view. First, does an energy barrier keep the particles apart? Second, which protein contacts form at the surface? Third, if a contact forms, does it survive motion in a realistic environment?",
            "I will use DLVO modeling, protein-protein docking, and coarse-grained molecular dynamics to address those questions in sequence."
        ],
        "cue": "This is the spine of the talk: barrier, contact, persistence. Repeat the same three words on Slide 12.",
    },
    {
        "n": 3,
        "title": "Why discuss nanoparticles in real foods?",
        "seconds": 85,
        "screen": "Bioactive compounds, a nanocarrier, the digestive tract, and three food matrices: milk, yogurt, and bread.",
        "script": [
            "Let me start with the practical goal. We would like to put compounds such as curcumin, vitamin D, or omega-3 into everyday foods. These bioactives can be poorly soluble, chemically unstable, or poorly absorbed on their own. A food-grade nanoparticle can act like a very small delivery vehicle: it protects the cargo, keeps it dispersed, and can improve bioavailability.",
            "But the vehicle has to survive the road. Many nanoparticle studies use a clean buffer or deionized water. A food matrix is much busier. It contains salts, changing pH, calcium, fat, sugars, and many competing proteins.",
            "As soon as the particle enters that environment, proteins adsorb onto its surface. That coating - the protein corona - can change the particle size, charge, aggregation behavior, and ultimately delivery performance. On the right are three simple target environments: milk at about pH 6.6 and 80 millimolar ionic strength, yogurt at pH 4.5 and 20 millimolar, and a bread-release condition near pH 5.3 and 50 millimolar."
        ],
        "cue": "Point from the payload to the carrier, then to the three foods. Pause briefly after 'the vehicle has to survive the road.'",
    },
    {
        "n": 4,
        "title": "Why model it computationally?",
        "seconds": 70,
        "screen": "A three-level workflow: DLVO, molecular docking, and coarse-grained MD.",
        "script": [
            "The experimental challenge is that screening is slow and material-intensive. Every new pH, salt level, or protein mixture means another preparation and another set of measurements. We can observe what happened, but it is harder to predict what we should try next.",
            "Our objective is to predict the carrier's behavior before we make every formulation. The workflow moves from coarse to fine. DLVO treats the particles as whole colloids and asks whether a stability barrier exists. Docking zooms down to amino-acid contacts and asks which interfaces are plausible. Coarse-grained MD then turns the static picture into a movie and asks whether those contacts persist.",
            "The methods are connected. The protein-layer repulsion estimated from MD can eventually feed back into an extended DLVO model."
        ],
        "cue": "Trace the three boxes from left to right, then point back to the first box when explaining the feedback loop.",
    },
    {
        "n": 5,
        "title": "The three proteins in play",
        "seconds": 80,
        "screen": "Structures and composition profiles for whey proteins, caseins, and alpha-zein.",
        "script": [
            "Before the modeling, I want to introduce the proteins, because they behave very differently.",
            "Whey proteins such as beta-lactoglobulin and alpha-lactalbumin are compact, folded, and water-soluble. Beta-lactoglobulin is about 18 kilodaltons and makes up roughly half of whey; alpha-lactalbumin is about 14 kilodaltons and contributes about 20 percent.",
            "Caseins are different. They do not have one fixed folded structure. The four major casein families assemble with calcium phosphate into micelles around 150 nanometers, and together they account for roughly 80 percent of milk protein.",
            "Zein is the least familiar one for many audiences. It is a maize storage protein, or prolamin. It is rich in hydrophobic amino acids - about 60 percent - and has only about 2 percent charged residues. That is why zein is poorly soluble in water and readily assembles into particles. The same chemistry that makes it a useful carrier also makes its surface structure difficult to define."
        ],
        "cue": "Keep this as an accessible comparison, not a protein-chemistry lecture. End on the structure problem to set up docking later.",
    },
    {
        "n": 6,
        "title": "DLVO: from DLS data to a stability barrier",
        "seconds": 85,
        "screen": "DLVO equations and an illustrative interaction-energy curve with the barrier height marked.",
        "script": [
            "The first model is classical DLVO theory. In plain language, it adds two competing interactions. Electrostatic double-layer repulsion pushes similarly charged particles apart, while van der Waals attraction pulls them together.",
            "When we add those terms, we obtain the black interaction-energy curve. The maximum is the energy barrier. You can think of aggregation as two particles trying to roll over a hill. A high hill slows them down; a small or absent hill allows rapid aggregation. Here we use about 15 kT as a stable threshold, 10 to 15 as marginal, and below 10 as unstable.",
            "The useful part is that most inputs are measurements we already collect. Particle radius and zeta potential come from DLS and ELS, while pH and ionic strength determine electrostatic screening. The Hamaker constant comes from the zein literature. The plotted curve is only a schematic to explain the method, not one of our measured conditions."
        ],
        "cue": "Point to the blue repulsion, red attraction, and then the black barrier. Say explicitly that the displayed curve is illustrative.",
    },
    {
        "n": 7,
        "title": "DLVO: measured grid and surface response",
        "seconds": 90,
        "screen": "Sample vials, zeta potential versus pH, and aggregation onset versus ionic strength.",
        "script": [
            "We then measured a three-by-three grid: pH 4, 5.5, and 7, crossed with ionic strengths of 10, 50, and 150 millimolar, with three replicates per condition.",
            "The center plot shows how the zeta potential changes with pH and salt. The fitted isoelectric point is around pH 4.94, where the net surface potential crosses zero. Bare zein is reported closer to about 6.2 in the comparison used here. That shift tells us the caseinate coating is not just floating nearby; it changes the chemistry the particle presents to solution.",
            "The right plot shows the physical consequence. At pH 4, the particle size stays near a few hundred nanometers at 10 and 50 millimolar, but jumps to about 8,000 nanometers at 150 millimolar. That is an aggregation signature, placing the critical coagulation window between 50 and 150 millimolar for this condition."
        ],
        "cue": "Move left to right: samples, charge, size. Avoid claiming a precise critical coagulation concentration; call it a window.",
    },
    {
        "n": 8,
        "title": "DLVO: findings and limitation",
        "seconds": 95,
        "screen": "A stability map with food targets and a ranked bar chart of measured and food-relevant energy barriers.",
        "script": [
            "Here is the main DLVO result. At 10 millimolar, the particles have a real barrier across the pH range: about 36 kT at pH 4, 17 at pH 5.5, and 49 at pH 7. So the formulation can be electrostatically stable under low-salt conditions.",
            "But the barrier collapses as ionic strength increases. At pH 7, it falls from 49 kT at 10 millimolar to 5.6 kT at 50 millimolar and essentially zero at 150 millimolar. Salt screens the surface charge, so the repulsive hill disappears.",
            "Now look at the stars on the stability map. Milk, yogurt, and bread-release all remain below the 10 kT threshold in classical DLVO. Yogurt comes closest at about 5.7 kT because it has the lowest ionic strength.",
            "This does not mean the real particles must instantly fail. Classical DLVO includes charge and van der Waals attraction, but it does not include steric repulsion from a protein coating or hydration forces. For protein-coated particles, those missing terms may be important. The disagreement between the simple model and a stable vial is therefore useful: it tells us which physics we still need to add."
        ],
        "cue": "The second point should be 50 mM, even though the slide bullet repeats 10 mM. Frame the missing steric term as a model boundary, not an excuse.",
    },
    {
        "n": 9,
        "title": "Docking: no protein wins",
        "seconds": 75,
        "screen": "The docking workflow and comparison of scores for five milk proteins.",
        "script": [
            "DLVO cannot tell us which protein makes contact, so we next used protein-protein docking. We started from predicted structures, restricted the search toward hydrophilic residues that could be exposed on zein, used MEGADOCK for a rigid-body search, and examined the top poses and contacts.",
            "We tested five milk proteins: two whey proteins and three caseins. The scores range only from 5.16 to 5.82. Although the bars have an order, the spread is too small and the structural uncertainty is too large to declare a meaningful winner.",
            "That distinction matters: docking ranks and generates hypotheses; it does not directly measure binding free energy. Here the defensible result is not that alpha-lactalbumin wins. The defensible result is that no protein is clearly distinguishable."
        ],
        "cue": "Do not read the bar ranking as affinity. Use the exact phrase: 'Docking ranks and hypothesizes; it does not measure.'",
    },
    {
        "n": 10,
        "title": "Docking: why no protein wins",
        "seconds": 100,
        "screen": "A zein-beta-lactoglobulin docking pose, five independent zein predictions, and model confidence coloring.",
        "script": [
            "Why can we not resolve selectivity? The interfaces from the restrained docking contain only three to six zein residues, and all five proteins converge around residues 37 to 39. A typical stable protein-protein interface is larger, so these contacts are too small to support a strong mechanistic ranking.",
            "The deeper problem is the receptor. Alpha-zein has no experimentally resolved, reproducible tertiary structure for this use. On the center image, five independent predictions of the same sequence produce very different global folds. In quantitative comparisons, the backbone RMSD ranges from roughly 14 to 22 angstroms. In other words, we are not docking against one reliable shape.",
            "The most confident segment, residues 84 to 115, is simply a regular alpha-helix. It aligns almost perfectly with a generic ideal helix, which confirms secondary structure but gives us no unique pocket or tertiary surface.",
            "So I view this as a negative result with a mechanism, not a failed experiment. Rigid one-to-one docking cannot represent a zein nanoparticle, where many chains create a crowded surface and a milk protein may contact several chains at once. The likely corona is non-specific and multivalent, so the question must move from a static single-chain pose to a dynamic surface patch."
        ],
        "cue": "Present the evidence chain in order: tiny interfaces, inconsistent folds, generic helix, then the multivalent interpretation.",
    },
    {
        "n": 11,
        "title": "Coarse-grained MD: the plan",
        "seconds": 105,
        "screen": "A whole 70 nm nanoparticle, a 12 nm surface patch, and a log-scale comparison of simulation size.",
        "script": [
            "The next level is coarse-grained molecular dynamics. Docking gives a still image; MD gives a movie. Coarse-graining groups roughly four heavy atoms into one interaction bead, which reduces the particle count and allows a larger time step.",
            "A whole 70-nanometer zein particle is still far too expensive. It contains roughly 6,000 zein chains, around 30 million protein atoms, and about 60 million atoms after adding water. Even coarse-grained, the full particle would require around 6 million beads.",
            "Instead, we simulate a local 12-nanometer surface patch made from about nine zein chains, with two to four milk proteins above it. This is not meant to be the whole particle. It is a local window where multivalent contacts can occur. Across that window, the curvature of a 70-nanometer particle differs from a flat surface by only about half a nanometer, so a flat patch is a reasonable first approximation.",
            "The patch contains about 20,000 beads - roughly 2,700 times fewer particles than the all-atom whole-particle system. We can then set pH-related protonation and ionic strength, follow whether the docking contacts persist, and estimate protein-layer repulsion. That effective steric term can be fed back into extended DLVO, closing the loop from Slide 4."
        ],
        "cue": "Say 'the next level' and 'plan' so the audience does not mistake this for completed data. Slow down for the half-nanometer curvature argument.",
    },
    {
        "n": 12,
        "title": "Summary and significance",
        "seconds": 80,
        "screen": "The three-tier pipeline repeated with the result or next output from each level.",
        "script": [
            "To summarize, the workflow asks three connected questions: barrier, contact, and persistence.",
            "First, DLVO shows that zein-caseinate nanoparticles can have a substantial electrostatic barrier at low ionic strength, but the barrier collapses with salt, and none of the three food conditions clears 10 kT in the classical model.",
            "Second, rigid docking cannot identify a selective winner among five milk proteins. We can explain why: the predicted zein fold is not reproducible, the interfaces are only three to six residues, and a real nanoparticle interaction is likely non-specific and multivalent.",
            "Third, coarse-grained MD will test those contacts on a multi-chain surface patch and provide the protein-layer repulsion that classical DLVO is missing.",
            "The broader impact is practical: fewer trial-and-error experiments, more rational nanocarrier design, and ultimately more reliable nutrient delivery in real foods."
        ],
        "cue": "Point to each box while repeating 'barrier, contact, persistence.' End on the practical impact, not on a limitation.",
    },
    {
        "n": 13,
        "title": "Acknowledgements",
        "seconds": 25,
        "screen": "USDA NIFA support, collaborator acknowledgement, contact information, and logos.",
        "script": [
            "I would like to thank USDA NIFA for supporting this work, Dr. Changqing Wu at the University of Delaware, my co-authors, and everyone in our group who contributed to the experimental and computational pieces.",
            "My contact information is shown here. Before questions, I have one brief announcement to share."
        ],
        "cue": "Do not read the grant number or email address aloud. Use the final sentence to make the transition to Slide 14 feel intentional.",
    },
    {
        "n": 14,
        "title": "Sustainable Food Technology announcement",
        "seconds": 20,
        "screen": "Journal scope, selected publication metrics, social links, and QR code.",
        "script": [
            "Sustainable Food Technology is a gold open-access journal from the Royal Society of Chemistry covering safe, high-quality, and environmentally sustainable food production. If that fits your work, the QR code and journal information are on the slide.",
            "Thank you very much. I am happy to take questions."
        ],
        "cue": "Keep this to two sentences. Do not read every metric; leave the slide visible for Q&A if required by the session.",
    },
]


QA = [
    (
        "Your particles are stable at 10 mM, but milk is around 80 mM. Will they aggregate in milk?",
        "Classical DLVO predicts that electrostatic stabilization alone is insufficient in milk-like ionic strength. I would state that directly. However, the model omits steric repulsion from the caseinate layer and hydration forces, so its barrier is not the complete interaction potential. The result tells us both where the formulation is weakest and which missing physics should be measured next."
    ),
    (
        "Why does the model predict little or no barrier when particles can still look stable in a vial?",
        "A visually stable vial can be supported by interactions that classical DLVO does not include, especially steric and hydration repulsion from adsorbed protein. It can also depend on observation time. I would not fit the Hamaker constant simply to force agreement; the mismatch is evidence that an extended model is needed."
    ),
    (
        "Why is the isoelectric point different from bare zein?",
        "The particle surface is zein plus a caseinate coating, so the chemistry exposed to water is not the chemistry of bare zein. The shift toward pH 4.94 is consistent with caseinate changing the surface. I would still describe 4.94 as a fitted estimate and collect finer pH spacing around the crossover before treating two decimal places as definitive."
    ),
    (
        "Why use an AlphaFold model if the confidence is so low?",
        "We do not treat it as a reliable receptor. We used the predictions to quantify the uncertainty: five models give very different global folds, with roughly 14 to 22 angstrom backbone RMSD. That evidence is part of the result. It shows why rigid docking cannot support a strong selectivity claim and why an ensemble or surface-patch method is more appropriate."
    ),
    (
        "Does the docking plot show that alpha-lactalbumin binds best?",
        "No. The numerical order is visible, but the score spread is small relative to the structural uncertainty, the interfaces are only three to six residues, and the receptor fold is not reproducible. I would report the proteins as indistinguishable, not rank them as measured affinities."
    ),
    (
        "How will you validate the computational predictions?",
        "The pipeline is intended to be paired with experiment. DLS size and zeta potential test the colloidal trends; SDS-PAGE or proteomics can identify which proteins remain in the hard corona; and spectroscopy can test structural changes. For MD, contact persistence and adsorption trends should be compared across the same pH and ionic-strength conditions used experimentally."
    ),
    (
        "Why simulate a flat patch instead of a whole nanoparticle?",
        "The whole system is computationally impractical, even after coarse-graining. A 12-nanometer window reduces the model to about 20,000 beads, and the curvature error across that window is only about 0.5 nanometer for a 70-nanometer particle. The patch is therefore a controlled local approximation, not a claim to reproduce the entire particle."
    ),
]


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)


def shade_paragraph(paragraph, fill=LIGHT, left_border=BLUE):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), left_border)
    borders.append(left)
    p_pr.append(borders)


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=9, color=GRAY)


def fmt_time(seconds):
    return f"{seconds // 60}:{seconds % 60:02d}"


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string("222222")
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in (
    ("Title", 28, NAVY, 0, 8),
    ("Subtitle", 14, GRAY, 0, 18),
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
):
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Remove the decorative bottom rule inherited by Word's built-in Title style.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

# Running header and footer.
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = hp.add_run("ACS Fall 2026  |  Speaker Script")
set_font(hr, size=9, color=GRAY, bold=True)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = fp.add_run("Yi Wang  |  ")
set_font(fr, size=9, color=GRAY)
add_field(fp, "PAGE")

# Cover page: workshop_agenda pattern.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(42)
p.paragraph_format.space_after = Pt(8)
r = p.add_run("ACS FALL 2026 - CHICAGO")
set_font(r, size=11, color=RED, bold=True)

p = doc.add_paragraph(style="Title")
p.add_run("Protein Corona Formation on Zein Nanoparticles in Dairy Protein Systems")
p = doc.add_paragraph(style="Subtitle")
p.add_run("Speaker script for a 15-17 minute oral presentation")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
r = p.add_run("Yi Wang, Liping Zhang, Qin Wang  |  Department of Nutrition and Food Science, University of Maryland")
set_font(r, size=10.5, color=GRAY)

metric = doc.add_table(rows=1, cols=4)
set_table_geometry(metric, [2340, 2340, 2340, 2340], indent_dxa=120)
metric_tr_pr = metric.rows[0]._tr.get_or_add_trPr()
metric_header = OxmlElement("w:tblHeader")
metric_header.set(qn("w:val"), "true")
metric_tr_pr.append(metric_header)
labels = [("TARGET", "16:15"), ("RANGE", "15-17 min"), ("DECK", "14 slides"), ("AUDIENCE", "Food/ag graduate")]
for cell, (label, value) in zip(metric.rows[0].cells, labels):
    set_cell_shading(cell, PALE_GOLD)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + "\n")
    set_font(r, size=8, color=RED, bold=True)
    r = p.add_run(value)
    set_font(r, size=11, color=NAVY, bold=True)

p = doc.add_paragraph(style="Heading 2")
p.add_run("Talk map")
p = doc.add_paragraph()
p.add_run("Context and proteins (Slides 1-5) -> stability barrier (Slides 6-8) -> molecular contacts (Slides 9-10) -> dynamic surface model (Slide 11) -> synthesis and close (Slides 12-14).")

p = doc.add_paragraph()
shade_paragraph(p, PALE_GOLD, RED)
r = p.add_run("Pre-talk slide check: ")
set_font(r, bold=True, color=RED)
r = p.add_run("On Slide 8, the middle salt condition in the third bullet should be 50 mM, not 10 mM. The graph is labeled correctly, and this script uses 50 mM.")
set_font(r, color=NAVY)

p = doc.add_paragraph()
shade_paragraph(p, LIGHT, BLUE)
r = p.add_run("Pacing options: ")
set_font(r, bold=True, color=BLUE)
r = p.add_run("For a 15-minute version, shorten Slides 5, 10, and 11 and keep Slide 14 to one sentence. For a 17-minute version, use the full text and pause briefly at the main figures.")

# One slide per page for podium usability.
running = 0
for slide in SLIDES:
    doc.add_page_break()
    running += slide["seconds"]
    h = doc.add_paragraph(style="Heading 1")
    h.paragraph_format.space_before = Pt(0)
    r = h.add_run(f"Slide {slide['n']} - {slide['title']}")
    set_font(r, size=16, color=BLUE, bold=True)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(6)
    r = meta.add_run(f"Target time {fmt_time(slide['seconds'])}  |  Running total {fmt_time(running)}")
    set_font(r, size=9.5, color=WHITE, bold=True)
    shade_paragraph(meta, RED, RED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("On screen: ")
    set_font(r, bold=True, color=DARK_BLUE)
    p.add_run(slide["screen"])

    h2 = doc.add_paragraph(style="Heading 2")
    h2.add_run("Script")
    for text in slide["script"]:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.22
        p.paragraph_format.space_after = Pt(7)
        p.add_run(text)

    p = doc.add_paragraph()
    shade_paragraph(p, LIGHT, BLUE)
    r = p.add_run("Delivery cue: ")
    set_font(r, bold=True, color=BLUE)
    r = p.add_run(slide["cue"])
    set_font(r, italic=True, color=GRAY)

# Q&A pages.
doc.add_page_break()
h = doc.add_paragraph(style="Heading 1")
h.paragraph_format.space_before = Pt(0)
h.add_run("Q&A preparation")
p = doc.add_paragraph()
p.add_run("Use the short answer first. Add detail only if the questioner asks a follow-up.")

for idx, (question, answer) in enumerate(QA, 1):
    if idx == 4:
        doc.add_page_break()
        h = doc.add_paragraph(style="Heading 1")
        h.paragraph_format.space_before = Pt(0)
        h.add_run("Q&A preparation - continued")
    q = doc.add_paragraph(style="Heading 2")
    q.add_run(question)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.add_run(answer)

doc.add_page_break()
h = doc.add_paragraph(style="Heading 1")
h.paragraph_format.space_before = Pt(0)
h.add_run("Phrases worth keeping exact")
for phrase in (
    "The particle does not keep the clean surface we designed; the corona gives it a new identity.",
    "Barrier, contact, and persistence.",
    "Docking ranks and hypothesizes; it does not measure.",
    "A negative result with a mechanism, not a failed experiment.",
    "Docking gives a still image; molecular dynamics gives a movie.",
    "The disagreement tells us which physics is missing.",
):
    p = doc.add_paragraph()
    shade_paragraph(p, LIGHT, GOLD)
    r = p.add_run(f'"{phrase}"')
    set_font(r, size=11.5, color=NAVY, italic=True)

# Update fields on open.
settings = doc.settings._element
update = OxmlElement("w:updateFields")
update.set(qn("w:val"), "true")
settings.append(update)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)

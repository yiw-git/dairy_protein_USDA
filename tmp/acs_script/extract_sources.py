from pathlib import Path
from docx import Document
from pptx import Presentation
import pdfplumber


ROOT = Path(r"C:\Users\nnjj1\UMD-work\dairy_protein_USDA")
ACS = ROOT / "report_summary" / "ACS2026"
OUT = ROOT / "tmp" / "acs_script"
OUT.mkdir(parents=True, exist_ok=True)


def iter_shape_text(shape, indent=""):
    lines = []
    if getattr(shape, "shape_type", None) == 6:  # GROUP
        for child in shape.shapes:
            lines.extend(iter_shape_text(child, indent + "  "))
        return lines
    if getattr(shape, "has_text_frame", False):
        text = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
        if text:
            lines.append(f"{indent}[TEXT] {text}")
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            vals = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            lines.append(f"{indent}[TABLE] " + " | ".join(vals))
    if getattr(shape, "has_chart", False):
        chart = shape.chart
        if chart.has_title:
            lines.append(f"{indent}[CHART TITLE] {chart.chart_title.text_frame.text}")
        for idx, series in enumerate(chart.series, 1):
            lines.append(f"{indent}[CHART SERIES {idx}] {series.name}")
    return lines


def extract_pptx(path):
    prs = Presentation(path)
    out = [f"FILE: {path}", f"SLIDES: {len(prs.slides)}"]
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"\n===== SLIDE {i} =====")
        for shape in slide.shapes:
            out.extend(iter_shape_text(shape))
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ""
        if notes:
            out.append("[NOTES]")
            out.append(notes)
    return "\n".join(out)


def extract_docx(path):
    doc = Document(path)
    out = [f"FILE: {path}"]
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            out.append(f"[{p.style.name}] {txt}")
    for ti, table in enumerate(doc.tables, 1):
        out.append(f"\n[TABLE {ti}]")
        for row in table.rows:
            out.append(" | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))
    return "\n".join(out)


def extract_pdf(path):
    out = [f"FILE: {path}"]
    with pdfplumber.open(path) as pdf:
        out.append(f"PAGES: {len(pdf.pages)}")
        for i, page in enumerate(pdf.pages, 1):
            out.append(f"\n===== PAGE {i} =====")
            out.append(page.extract_text(x_tolerance=1, y_tolerance=3) or "")
    return "\n".join(out)


(OUT / "slides.txt").write_text(
    extract_pptx(ACS / "ACS2026_slides_YW_0821.pptx"), encoding="utf-8"
)
(OUT / "speaker_script_existing.txt").write_text(
    extract_docx(ACS / "ACS2026_speaker_script.docx"), encoding="utf-8"
)
(OUT / "acs_scripts_existing.txt").write_text(
    extract_docx(ACS / "ACS-Scripts.docx"), encoding="utf-8"
)
(OUT / "slides_pdf.txt").write_text(
    extract_pdf(ACS / "ACS2026_slides_YW_0821.pdf"), encoding="utf-8"
)
